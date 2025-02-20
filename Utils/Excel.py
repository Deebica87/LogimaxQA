import pandas as pd
import re
import win32com.client
from openpyxl import load_workbook
from PIL import ImageGrab
import random
import driver



class ExcelUtils:
    file_path = "D:\CRM\Taneira\log.xlsx"
    excel_app = win32com.client.Dispatch("Excel.Application")
    excel_app.Quit()
    # if excel_app.Workbooks.Count > 0:
    #         for workbook in excel_app.Workbooks:
    #             workbook.Save()  # Save the workbook
    #             workbook.Close()
    
    # for workbook in excel_app.Workbooks:
    #     workbook.Save()  # Save the workbook
    #     workbook.Close()

    # open Excel file
    def read_excel(file_path):
        df = pd.read_excel(file_path)
        return df.fillna("")
    
    #To get sheet names in a list from workbook
    def get_sheet_names(file_path):
        excel_file = pd.ExcelFile(file_path)
        File = excel_file.sheet_names
        noOfSheets = len(excel_file.sheet_names)
        sheet_names = list(File)
        #print(noOfSheets)
        return sheet_names

    #Read & Get value from master sheet data
    def get_master_sheet_data(file_path):
        """
        Reads the 'Master' sheet and returns the functions marked for execution.
        """
        df = pd.read_excel(file_path, sheet_name="Master")
        print(df)
        functions_to_execute = []
        for index, row in df.iterrows():
            if re.match(r"yes", str(row["Execution"]), re.IGNORECASE):
                functions_to_execute.append(row["Function"])
                print(functions_to_execute)
        return functions_to_execute
    
    #To Fetch Valid Rows
    def get_valid_rows (file_path,function_name):
        workbook = load_workbook(file_path)
            # workbook = load_workbook(file_path)
        sheet = workbook[function_name]
        i =2
        count = 0
        while (i<100):
            cellvalue = sheet.cell(row =i, column=1).value
            if cellvalue is None:
                break
            i=i+1
            count = count+1
        return(count+2)
    
    #Fetch No Of Pass & Fail
    def get_Status(file_path,function_name):
        workbook = load_workbook(file_path)
        sheet = workbook[function_name]
        count = (ExcelUtils.get_valid_rows(file_path,function_name))+1
        Pass = 0
        Fail = 0
        i = 2
        while(i<=count):
            cellvalue =sheet.cell(row =i, column=2).value
            if cellvalue and cellvalue.strip().lower() == "pass":  # Case-insensitive comparison
               Pass += 1
            elif cellvalue and cellvalue.strip().lower() == "fail":  # Case-insensitive comparison
               Fail += 1
            # if cellvalue == "Pass":
            #     Pass = Pass+1
            # elif cellvalue == "Fail":
            #     Fail = Fail+1   
            i=i+1    
        status = (f"Pass {Pass}, Fail {Fail}")
        print(status)
        return(status)

    #Update Master Status
    def update_master_status(file_path,Status,function_name):
        
        workbook = load_workbook(file_path)
            # workbook = load_workbook(file_path)
        sheet = workbook["Master"]
        print(Status)
        i =2
        while(i<=100):
            cellvalue =sheet.cell(row =i, column=1).value
            if cellvalue == function_name:
                sheet.cell(row=i,column=3).value = Status 
                workbook.save(file_path)
                break
            i=i+1 
            
    #Get screenshot
    def screenshot(function_name):
        path = f"D:\\CRM\\Taneira\\screeshot\\{function_name}_{random.random():.4f}.png"
        driver.save_screenshot(path)
        print(path)
        
                 
                
        
        
                  
            
        