from selenium import webdriver
from selenium.webdriver.common.by import By
from time import sleep
from Utils.Excel import ExcelUtils
from openpyxl import load_workbook
import allure
import random
from PIL import ImageGrab

FILE_PATH = ExcelUtils.file_path
class LoginAutomation:
    def __init__(self, driver):
        self.driver = driver
    def perform_login(self):
        try:
            function_name = "Login"
            valid_rows = ExcelUtils.get_valid_rows(FILE_PATH,function_name)
            workbook = load_workbook(FILE_PATH)
            sheet = workbook[function_name]
            for row_num in range(2,valid_rows):   
                test_case_id = sheet.cell(row=row_num, column=1).value
                Expected_status=sheet.cell(row=row_num, column=4).value
                print(Expected_status)
                url = sheet.cell(row=row_num, column=5).value
                print(url)
                username = sheet.cell(row=row_num, column=6).value
                password = sheet.cell(row=row_num, column=7).value  
                 
                self.driver.get(url)
                #Enter username
                username_field = self.driver.find_element(By.NAME, "username")   #Update locator if required
                username_field.clear()
                username_field.send_keys(username)

                #Enter password
                password_field = self.driver.find_element(By.NAME, "password")   #Update locator if required
                password_field.clear()
                password_field.send_keys(password)

                #Click Login
                login_button = self.driver.find_element(By.XPATH, "//button[@type='submit']")   #Update locator
                login_button.click()
                # path = "D:\CRM\Taneira\screeshot\login"+str(random.random())+".jpg"
                # self.driver.save_screenshot(path)
                ExcelUtils.screenshot(function_name)
                sleep(5)
                match Expected_status:
                        case "valid Credentials":                            
                            if self.driver.find_element(By.XPATH, '//h1[contains(text(), "Dashboard")]'):
                                test_status = "Pass"
                                Actual_status = "Login successful"
                            else:
                                test_status = "Fail"
                                Actual_status = "Login unsuccessful"
                        case "Invalid Credentials":
                            if self.driver.find_element(By.XPATH, '//h1[contains(text(), "Dashboard")]'):
                                test_status = "Fail"
                                Actual_status = "Login successful"
                            else:    
                                test_status = "Pass"
                                Actual_status = "Login Unsuccessful" 
                sheet.cell(row=row_num, column=2).value = test_status  # Write Test Status
                sheet.cell(row=row_num, column=3).value = Actual_status  # Write Actual Status
                workbook.save(FILE_PATH)
                if valid_rows-2>1:
                    print(valid_rows)
                    sleep(2)
                    try:
                        if (self.driver.find_element(By.PARTIAL_LINK_TEXT, "Lmxsupport")):
                            self.driver.find_element(By.PARTIAL_LINK_TEXT, "Lmxsupport").click()
                            self.driver.find_element(By.PARTIAL_LINK_TEXT, "Sign out").click()  
                    except Exception as e:   
                        print (f"Error during login: {e}" )      
            print("Login Funtion Completed")   
            Status = ExcelUtils.get_Status(FILE_PATH,function_name)  
            print(Status)
            Update_master = ExcelUtils.update_master_status(FILE_PATH,Status,function_name)
                                    
        except Exception as e:
            print(f"Error during login: {e}" )
            
    def url(self):
        try:
            function_name = "Login"
            valid_rows = ExcelUtils.get_valid_rows(FILE_PATH,function_name)
            workbook = load_workbook(FILE_PATH)
            sheet = workbook[function_name]
            for row_num in range(2,valid_rows):   
                test_case_id = sheet.cell(row=row_num, column=1).value
                Expected_status=sheet.cell(row=row_num, column=4).value
                print(Expected_status)
                url = sheet.cell(row=row_num, column=5).value
                return url                       
        except Exception as e:
            print(f"Error during login: {e}" )
    